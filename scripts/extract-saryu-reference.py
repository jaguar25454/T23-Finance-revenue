"""Read-only source extraction. The workbook supplies classification, not sales history."""
import csv, hashlib, json, pathlib, sys
import openpyxl

folder = pathlib.Path(sys.argv[1])
target = pathlib.Path(__file__).resolve().parents[1] / 'app/data/saryu-reference.json'
book = folder / 'Kinn1_Daily_Sales_Report_260823.xlsx'
w = openpyxl.load_workbook(book, data_only=False)
# Sale Kinn1 C:N: row 10 is Type, row 11 is Sub Type, rows 8/9 are SUMIFS criteria.
groups = []
for col, kind in [(3,'Food'),(4,'Food'),(6,'Beverage - Non-Alcohol'),(7,'Beverage - Non-Alcohol'),(8,'Beverage - Non-Alcohol'),(10,'Beverage - Alcohol'),(11,'Beverage - Alcohol'),(12,'Beverage - Alcohol'),(14,'Other')]:
    s=w['Sale Kinn1']
    groups.append(dict(subgroup=str(s.cell(11,col).value or s.cell(9,col).value).strip(),type=kind,category=kind,criteria=[str(s.cell(row,col).value).strip() for row in (8,9) if s.cell(row,col).value],source=f'Sale Kinn1 {s.cell(11,col).coordinate}'))
def clean(v): return ' '.join(str(v or '').split())
def key(name, group, category): return json.dumps([clean(name).lower(),clean(group).lower(),clean(category).lower()],ensure_ascii=False,separators=(',',':'))
catalog={}
for row in list(w['Data Sale'].values)[1:]:
    if not row[6] or row[0] in ('Total',None): continue
    name,group,category=map(clean,(row[6],row[24],row[25]))
    identity=key(name,group,category)
    entry=catalog.setdefault(identity,dict(id=identity,name=name,posGroup=group,sourceCategory=category,prices=[],source='Data Sale'))
    if isinstance(row[9],(float,int)) and row[9] not in entry['prices']: entry['prices'].append(row[9])
files=[]
for file in sorted(folder.glob('*.csv')):
    content=file.read_text(encoding='utf-8-sig')
    files.append(dict(name=file.name,sha256=hashlib.sha256(file.read_bytes()).hexdigest(),text=content))
result=dict(workbook=dict(name=book.name,sha256=hashlib.sha256(book.read_bytes()).hexdigest()),groups=groups,catalog=list(catalog.values()),files=files)
target.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(catalog=len(catalog),groups=groups,files=[f['name'] for f in files]),ensure_ascii=False))
