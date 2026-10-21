"""Read-only extraction of the supplied Finance reference; never saves the XLSX.

Usage: python scripts/extract-capricci-finance.py <reference.xlsx>
Generated JSON keeps source coordinates and cached checks separate from inputs.
"""
import collections
import hashlib
import json
from pathlib import Path
import sys

import openpyxl

source = Path(sys.argv[1])
before = hashlib.sha256(source.read_bytes()).hexdigest()
values = openpyxl.load_workbook(source, read_only=False, data_only=True)
formulas = openpyxl.load_workbook(source, read_only=False, data_only=False)
category = values['Category']
# D contains vertical merged category labels. Only inherit within an actual
# merged range; later blank rows must not inherit the final DOLCI category.
category_cells = {}
for merged in category.merged_cells.ranges:
    if merged.min_col == 4 and merged.max_col == 4:
        for row_number in range(merged.min_row, merged.max_row + 1):
            category_cells[row_number] = category.cell(merged.min_row, 4)
mapping = []
for row in category.iter_rows(min_row=4):
    code = row[5].value
    if not code or str(code).lower() == 'code':
        continue
    category_cell = category_cells.get(row[5].row, row[3])
    mapping.append({
        'sourceRow': row[5].row, 'itemCode': str(code),
        'subgroup': row[10].value, 'menu': row[6].value,
        'sourceCategory': category_cell.value,
        'categorySourceCell': category_cell.coordinate if category_cell.value else None,
    })
groups = [
    {'sourceRow': row[14].row, 'subgroup': row[14].value,
     'mainGroup': row[15].value, 'category': row[16].value}
    for row in category.iter_rows(min_row=4, max_row=21)
]
rows = []
for row in values['S104'].iter_rows(min_row=2):
    if row[7].value is None:
        continue
    i = row[7].row
    rows.append({
        'sourceRow': i, 'businessDate': row[17].value.strftime('%Y-%m-%d'),
        'customerName': row[3].value, 'saleId': row[4].value, 'saleNum': row[6].value,
        'itemCode': row[7].value, 'itemDesc': row[8].value, 'qty': row[10].value,
        'extendedPrice': row[11].value, 'adjustment': row[15].value or 0,
        'cachedTotal': row[16].value, 'cachedSubgroup': row[18].value,
        'cachedMainGroup': row[19].value, 'cachedCategory': row[20].value,
    })
dates = sorted(set(row['businessDate'] for row in rows))
assert dates == ['2026-09-01'], dates
codes = collections.defaultdict(list)
for row in mapping:
    codes[row['itemCode']].append(row)
duplicate_conflicts = [
    {'itemCode': code, 'rows': [r['sourceRow'] for r in entries],
     'subgroups': [r['subgroup'] for r in entries]}
    for code, entries in codes.items()
    if len(set(r['subgroup'] for r in entries)) > 1
]
checks = []
for sheet, addresses in {
    'DashBord': ['J2', 'B3', 'E3', 'E4', 'E5', 'E6', 'G3', 'J3', 'J4', 'J5', 'J6'],
    'Sale Capricci ': ['C12', 'D12', 'E12', 'F12', 'G12', 'H12', 'I12', 'J12', 'K12', 'L12', 'M12', 'N12', 'O12', 'P12', 'Q12', 'R12', 'S12', 'T12', 'U12', 'V12', 'W12', 'Y12', 'Z12', 'Z43'],
    'S104': ['Q2', 'S2', 'T2', 'U2'],
}.items():
    for address in addresses:
        checks.append({'sheet': sheet, 'cell': address,
                       'formula': formulas[sheet][address].value,
                       'cachedValue': values[sheet][address].value})
payload = {
    'source': {'filename': source.name, 'sha256': before, 'businessDate': dates[0],
               'mappingRange': 'Category!F4:L897', 'groupRange': 'Category!O4:Q21',
               'itemCategoryRange': 'Category!D4:D897',
               'transactionRange': 'S104!A2:U134'},
    'customerFilter': 'ลูกค้าทั่วไป เงินสด',
    'categoryNames': ['Food', 'Beverage - Non-Alcohol', 'Beverage - Alcohol', 'Other'],
    'mapping': mapping, 'groups': groups, 'rows': rows,
    'duplicateConflicts': duplicate_conflicts, 'referenceChecks': checks,
}
assert before == hashlib.sha256(source.read_bytes()).hexdigest(), 'Original changed'
output = Path(__file__).resolve().parents[1] / 'app/data/capricci-finance-20260901.json'
output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
print(json.dumps({'rows': len(rows), 'mappingRows': len(mapping),
                  'duplicateConflicts': duplicate_conflicts, 'sha256': before}, ensure_ascii=True))
