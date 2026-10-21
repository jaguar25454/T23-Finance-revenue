"""Build the checked-in Maison reference snapshot from read-only source files."""
from __future__ import annotations

import csv
import hashlib
import json
import pathlib
import sys

import openpyxl


folder = pathlib.Path(sys.argv[1])
target = pathlib.Path(__file__).resolve().parents[1] / "app/data/maison-reference.json"
book = folder / "Alpea_Daily_Sales_Report_260822-For Codex.xlsx"
workbook = openpyxl.load_workbook(book, data_only=True)


def clean(value: object) -> str:
    return " ".join(str(value or "").split())


def identity(name: str, group: str, category: str) -> str:
    return json.dumps([clean(name).lower(), clean(group).lower(), clean(category).lower()], ensure_ascii=False, separators=(",", ":"))


# Finance's Sale Alpea structure. Pop Up is applied per transaction from A02 invoices,
# not assigned permanently to a menu in Master Data.
group_types = {
    "Set Menu": "Food", "A La Carte": "Food", "Other Food": "Food",
    "Water": "Beverage - Non-Alcohol", "Soft Drink": "Beverage - Non-Alcohol",
    "Mocktail": "Beverage - Non-Alcohol", "Coffee/Tea": "Beverage - Non-Alcohol",
    "Wine Glass": "Beverage - Alcohol", "Wine Bottle": "Beverage - Alcohol",
    "Cocktail": "Beverage - Alcohol", "Beers": "Beverage - Alcohol",
    "Hard Liquor": "Beverage - Alcohol", "Corkage Charge": "Beverage - Alcohol",
    "Other": "Other", "Others": "Other",
}

criteria: dict[str, list[str]] = {name: [] for name in group_types}
for product, _group, report_group in workbook["Code"].iter_rows(min_row=2, values_only=True):
    product_name, subgroup = clean(product), clean(report_group)
    if product_name and subgroup in criteria and product_name not in criteria[subgroup]:
        criteria[subgroup].append(product_name)
groups = [
    {"subgroup": subgroup, "type": output, "category": output, "criteria": values, "source": "Code A:C + Sale Alpea rows 10-11"}
    for subgroup, output in group_types.items()
    for values in [criteria[subgroup]]
    if values or subgroup == "Other"
]

catalog: dict[str, dict[str, object]] = {}
for row in workbook["Data Sale"].iter_rows(min_row=2, values_only=True):
    if not row[6] or row[0] in ("Total", None):
        continue
    name, group, category = map(clean, (row[6], row[24], row[25]))
    item_id = identity(name, group, category)
    entry = catalog.setdefault(item_id, {
        "id": item_id, "name": name, "posGroup": group, "sourceCategory": category,
        "prices": [], "source": "Data Sale",
    })
    if isinstance(row[9], (float, int)) and row[9] not in entry["prices"]:
        entry["prices"].append(row[9])

files = []
for file in sorted(folder.glob("*.csv")):
    text = file.read_text(encoding="utf-8-sig")
    files.append({"name": file.name, "sha256": hashlib.sha256(file.read_bytes()).hexdigest(), "text": text})

result = {
    "workbook": {"name": book.name, "sha256": hashlib.sha256(book.read_bytes()).hexdigest()},
    "branch": "Maison Dunand and Alpea",
    "groups": groups,
    "catalog": list(catalog.values()),
    "files": files,
}
target.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"catalog": len(catalog), "groups": groups, "files": [file["name"] for file in files]}, ensure_ascii=False))
